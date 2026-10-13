"""Extracao de texto dos documentos enviados.

`extract_text(path, ext)` devolve `(texto, ok)`. Qualquer falha de parsing e
capturada: o upload continua, apenas sem conteudo indexado (ok=False).
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_CHARS = 2_000_000  # limite defensivo para nao estourar o banco


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_spreadsheet(ext: str):
    def handler(path: Path) -> str:
        from .sheets import read_sheets, sheets_to_text

        return sheets_to_text(read_sheets(path, ext))

    return handler


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text:
                        parts.append(text)
    return "\n".join(parts)


_HANDLERS = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "xlsx": _extract_spreadsheet("xlsx"),
    "xlsm": _extract_spreadsheet("xlsm"),
    "xls": _extract_spreadsheet("xls"),
    "pptx": _extract_pptx,
}

_PLAIN = {"txt", "md", "markdown", "csv"}


def extract_text(path, ext: str):
    """Retorna (texto, ok). ext sem ponto, minusculo."""
    path = Path(path)
    ext = ext.lower().lstrip(".")
    try:
        if ext in _PLAIN:
            text = _read_text_file(path)
        elif ext in _HANDLERS:
            text = _HANDLERS[ext](path)
        else:
            logger.info("Sem extrator para .%s (%s)", ext, path.name)
            return "", False
        text = (text or "").strip()
        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS]
        return text, bool(text)
    except Exception:  # noqa: BLE001 - best effort, nao pode quebrar o upload
        logger.exception("Falha ao extrair texto de %s", path)
        return "", False


def _clean(value) -> str:
    return str(value).strip() if value else ""


def _meta_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    info = PdfReader(str(path)).metadata or {}
    return {"title": _clean(info.title), "summary": _clean(info.subject)}


def _meta_docx(path: Path) -> dict:
    from docx import Document as DocxDocument

    cp = DocxDocument(str(path)).core_properties
    return {
        "title": _clean(cp.title),
        "summary": _clean(cp.subject) or _clean(cp.comments),
    }


def _meta_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True)
    props = wb.properties
    result = {"title": _clean(props.title), "summary": _clean(props.description)}
    wb.close()
    return result


def _meta_pptx(path: Path) -> dict:
    from pptx import Presentation

    cp = Presentation(str(path)).core_properties
    return {
        "title": _clean(cp.title),
        "summary": _clean(cp.subject) or _clean(cp.comments),
    }


_META_HANDLERS = {
    "pdf": _meta_pdf,
    "docx": _meta_docx,
    "xlsx": _meta_xlsx,
    "xlsm": _meta_xlsx,
    "pptx": _meta_pptx,
}


def extract_doc_meta(path, ext: str) -> dict:
    """Le titulo/resumo embutidos no arquivo. Sempre devolve {'title','summary'}."""
    ext = ext.lower().lstrip(".")
    handler = _META_HANDLERS.get(ext)
    if not handler:
        return {"title": "", "summary": ""}
    try:
        data = handler(Path(path))
        return {"title": data.get("title", ""), "summary": data.get("summary", "")}
    except Exception:  # noqa: BLE001
        logger.exception("Falha ao ler metadados de %s", path)
        return {"title": "", "summary": ""}
