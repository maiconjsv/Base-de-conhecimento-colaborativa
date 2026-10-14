"""Planilhas (XLSX/XLSM/XLS/CSV) lidas de forma estruturada.

Em vez de virar so texto corrido, cada planilha alimenta o banco com:
- `DocumentSheet`: uma por aba (nome, cabecalho detectado, colunas, contagem);
- `SheetRow`: uma por linha nao vazia, com o NUMERO REAL da linha no Excel e as
  celulas ja formatadas como texto (datas dd/mm/aaaa, decimais com virgula).

Isso permite mostrar a "pseudoplanilha" no sistema e, na busca/assistente,
apontar a fonte exata: arquivo > aba > linha.

O texto indexado no FTS (`Document.content`) e gerado a partir das mesmas
linhas, por `sheets_to_text`.
"""

import csv
import io
import json
import logging
from datetime import date, datetime, time
from pathlib import Path

from flask import render_template, url_for
from markupsafe import Markup, escape
from sqlalchemy import insert, literal_column, select, table, text

from .db import db
from .locator import render_window

logger = logging.getLogger(__name__)

SPREADSHEET_EXTS = {"xlsx", "xlsm", "xls", "csv"}

MAX_SHEETS = 50
MAX_ROWS_PER_SHEET = 20_000
MAX_COLS = 80
MAX_CELL_CHARS = 1_000
# Ate este tamanho o XLSX e aberto por inteiro, o que da acesso a links e celulas
# mescladas; acima, modo "somente leitura" (mais leve, so os valores).
FULL_LOAD_MAX_BYTES = 25 * 1024 * 1024
MAX_MERGE_FILL = 200_000
_HEADER_SCAN_ROWS = 10


def is_spreadsheet(ext: str) -> bool:
    return (ext or "").lower().lstrip(".") in SPREADSHEET_EXTS


def column_letter(index: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA."""
    out = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


# --------------------------------------------------------------------------- #
# Formatacao de valores
# --------------------------------------------------------------------------- #
def format_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%d/%m/%Y %H:%M")
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.10g}".replace(".", ",")
    text = " ".join(str(value).split())
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS] + "…"
    return text


def _is_texty(value) -> bool:
    """Celula com cara de rotulo (texto que nao e numero)."""
    if not isinstance(value, str):
        return False
    probe = value.strip().replace(".", "").replace(",", "").replace("-", "")
    return bool(probe) and not probe.isdigit()


# --------------------------------------------------------------------------- #
# Leitura dos formatos -> [(nome_aba, [(numero_linha, [valores crus])])]
# --------------------------------------------------------------------------- #
def _has_content(values) -> bool:
    return any(v is not None and str(v).strip() for v in values)


def _with_link(value, cell):
    """Texto da celula + destino do link, quando o texto visivel nao o mostra
    (ex.: "Scalation Vero" -> https://...). Assim o endereco tambem e pesquisavel."""
    link = getattr(cell, "hyperlink", None)
    target = (getattr(link, "target", None) or "").strip() if link else ""
    if not target:
        return value
    shown = target[7:] if target.lower().startswith("mailto:") else target
    if value is None or not str(value).strip():
        return shown
    if shown.lower() in str(value).lower():
        return value
    return f"{value} ({shown})"


def _vertical_merges(ws) -> dict:
    """{(linha, coluna): valor} para as celulas cobertas por mesclagens verticais.

    No Excel, uma celula mesclada por varias linhas (ex.: o nome da operadora ao
    lado dos contatos dela) so guarda o valor na primeira. Repetir o valor nas
    linhas de baixo faz cada linha encontrada na busca dizer a quem pertence.
    Mesclagens so horizontais (titulos) nao sao repetidas.
    """
    fill = {}
    for rng in ws.merged_cells.ranges:
        if rng.max_row == rng.min_row:
            continue
        value = ws.cell(rng.min_row, rng.min_col).value
        if value is None:
            continue
        for r in range(rng.min_row + 1, rng.max_row + 1):
            fill[(r, rng.min_col)] = value
            if len(fill) >= MAX_MERGE_FILL:
                return fill
    return fill


def _raw_xlsx(path: Path):
    from openpyxl import load_workbook

    full = path.stat().st_size <= FULL_LOAD_MAX_BYTES
    wb = load_workbook(str(path), read_only=not full, data_only=True)
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows = []
            if full:
                merged = _vertical_merges(ws)
                # min_row/min_col=1: a posicao na lista = linha/coluna reais do Excel
                for number, cells in enumerate(
                    ws.iter_rows(min_row=1, min_col=1, max_col=min(ws.max_column, MAX_COLS)), start=1
                ):
                    values = []
                    for c, cell in enumerate(cells, start=1):
                        value = _with_link(cell.value, cell)
                        values.append(merged.get((number, c)) if value is None else value)
                    if _has_content(values):
                        rows.append((number, values))
                        if len(rows) >= MAX_ROWS_PER_SHEET:
                            break
            else:
                ws.reset_dimensions()  # dimensao gravada no arquivo pode estar errada
                # iter_rows comeca na linha 1 e devolve linhas vazias nos buracos,
                # entao a posicao = numero real da linha no Excel.
                for number, values in enumerate(ws.iter_rows(values_only=True), start=1):
                    if _has_content(values):
                        rows.append((number, list(values)))
                        if len(rows) >= MAX_ROWS_PER_SHEET:
                            break
            yield ws.title, rows
    finally:
        wb.close()


def _raw_xls(path: Path):
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        for idx in range(min(book.nsheets, MAX_SHEETS)):
            sh = book.sheet_by_index(idx)
            rows = []
            for r in range(sh.nrows):
                values = []
                for c in range(min(sh.ncols, MAX_COLS)):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        except Exception:  # noqa: BLE001 - data invalida
                            values.append(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        values.append(bool(cell.value))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        values.append(None)
                    else:
                        values.append(cell.value)
                if any(v is not None and str(v).strip() for v in values):
                    rows.append((r + 1, values))
                    if len(rows) >= MAX_ROWS_PER_SHEET:
                        break
            yield sh.name, rows
            book.unload_sheet(idx)
    finally:
        book.release_resources()


def _raw_csv(path: Path):
    from .extract import _read_text_file

    text = _read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for number, values in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        if any(v.strip() for v in values):
            rows.append((number, values))
            if len(rows) >= MAX_ROWS_PER_SHEET:
                break
    yield "Dados", rows  # o arquivo em disco tem nome aleatorio; CSV tem uma "aba" so


_READERS = {"xlsx": _raw_xlsx, "xlsm": _raw_xlsx, "xls": _raw_xls, "csv": _raw_csv}


# --------------------------------------------------------------------------- #
# Estruturacao: recorte de colunas + cabecalho
# --------------------------------------------------------------------------- #
def _detect_header(rows, width):
    """Indice (em `rows`) da linha de cabecalho, ou None.

    Cabecalho = uma das primeiras linhas, com >= 2 celulas preenchidas, todas
    com cara de rotulo, cobrindo boa parte das colunas, seguida de dados.
    """
    for i, (_, values) in enumerate(rows[:_HEADER_SCAN_ROWS]):
        filled = [v for v in values if v is not None and str(v).strip()]
        if len(filled) < 2 or len(filled) < width * 0.5:
            continue
        if all(_is_texty(v) for v in filled) and i + 1 < len(rows):
            return i
        # primeira linha "larga" que nao e cabecalho: os dados ja comecaram
        return None
    return None


def _structure(name, raw_rows):
    """Uma aba crua -> dict pronto para gravar."""
    first_col, last_col = None, -1
    for _, values in raw_rows:
        for c, v in enumerate(values[:MAX_COLS]):
            if v is not None and str(v).strip():
                first_col = c if first_col is None else min(first_col, c)
                last_col = max(last_col, c)
    if first_col is None:
        return {"name": name, "col_offset": 0, "header": [], "header_row": 0, "rows": []}

    width = last_col - first_col + 1

    def cut(values):
        vals = list(values[first_col:last_col + 1])
        vals += [None] * (width - len(vals))
        return vals

    trimmed = [(n, cut(v)) for n, v in raw_rows]
    h = _detect_header(trimmed, width)
    if h is not None:
        header_row = trimmed[h][0]
        header = [format_value(v) for v in trimmed[h][1]]
        body = trimmed[h + 1:]
    else:
        header_row, header, body = 0, [], trimmed
    # coluna sem nome no cabecalho recebe a letra do Excel
    header = [hv or column_letter(first_col + i) for i, hv in enumerate(header)] if header else []

    rows = [(n, [format_value(v) for v in vals]) for n, vals in body]
    return {
        "name": name,
        "col_offset": first_col,
        "header": header,
        "header_row": header_row,
        "rows": rows,
    }


def read_sheets(path, ext: str) -> list[dict]:
    """Le a planilha inteira. Pode levantar excecao (arquivo corrompido etc.)."""
    reader = _READERS[ext.lower().lstrip(".")]
    return [_structure(name, rows) for name, rows in reader(Path(path))]


def sheets_to_text(sheets: list[dict]) -> str:
    """Texto para o indice full-text: aba, cabecalho e linhas."""
    parts = []
    for sh in sheets:
        parts.append(f"# {sh['name']}")
        if sh["header"]:
            parts.append(" | ".join(sh["header"]))
        for _, cells in sh["rows"]:
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


def describe_sheets(sheets: list[dict], limit: int = 320) -> str:
    """Resumo legivel: abas, quantidade de linhas e colunas."""
    parts = []
    for sh in sheets:
        if not sh["rows"]:
            continue
        n = len(sh["rows"])
        piece = f"{sh['name']} ({n} {'linha' if n == 1 else 'linhas'}"
        if sh["header"]:
            cols = ", ".join(sh["header"][:6]) + ("…" if len(sh["header"]) > 6 else "")
            piece += f"; colunas {cols}"
        parts.append(piece + ")")
    if not parts:
        return "Planilha sem dados."
    label = "Planilha com 1 aba" if len(parts) == 1 else f"Planilha com {len(parts)} abas"
    out = f"{label}: " + "; ".join(parts) + "."
    if len(out) > limit:
        out = out[:limit].rsplit(" ", 1)[0] + "…"
    return out


def extract_spreadsheet(path, ext: str):
    """(texto, ok, sheets). Best effort: falha devolve ("", False, None)."""
    try:
        sheets = read_sheets(path, ext)
    except Exception:  # noqa: BLE001 - nao pode quebrar o upload
        logger.exception("Falha ao ler planilha %s", path)
        return "", False, None
    text = sheets_to_text(sheets).strip()
    return text, bool(text), sheets


# --------------------------------------------------------------------------- #
# Banco
# --------------------------------------------------------------------------- #
def store_sheets(doc, sheets: list[dict]) -> None:
    """Grava abas e linhas do documento (substitui as anteriores). Sem commit."""
    from .models import DocumentSheet, SheetRow

    clear_sheets(doc)
    for ordinal, sh in enumerate(sheets):
        n_cols = len(sh["header"]) or max((len(c) for _, c in sh["rows"]), default=0)
        sheet = DocumentSheet(
            document_id=doc.id,
            ordinal=ordinal,
            name=sh["name"][:255] or f"Aba {ordinal + 1}",
            header=json.dumps(sh["header"], ensure_ascii=False),
            header_row=sh["header_row"],
            col_offset=sh["col_offset"],
            n_cols=n_cols,
            n_rows=len(sh["rows"]),
            truncated=len(sh["rows"]) >= MAX_ROWS_PER_SHEET,
        )
        db.session.add(sheet)
        db.session.flush()
        if sh["rows"]:
            db.session.execute(
                insert(SheetRow),
                [
                    {
                        "sheet_id": sheet.id,
                        "document_id": doc.id,
                        "row_number": number,
                        "cells": json.dumps(cells, ensure_ascii=False),
                        "text": " | ".join(c for c in cells if c),
                    }
                    for number, cells in sh["rows"]
                ],
            )


def clear_sheets(doc) -> None:
    from .models import DocumentSheet, SheetRow

    SheetRow.query.filter_by(document_id=doc.id).delete(synchronize_session=False)
    DocumentSheet.query.filter_by(document_id=doc.id).delete(synchronize_session=False)
    db.session.expire(doc, ["sheets"])


def backfill_sheets() -> int:
    """Estrutura planilhas enviadas antes deste recurso existir.

    Roda na inicializacao; so pega planilhas que ainda nao tem nenhuma aba no
    banco. Devolve quantas foram processadas.
    """
    from flask import current_app

    from .models import Document, DocumentSheet

    has_sheets = db.session.query(DocumentSheet.document_id).distinct()
    pending = [
        d for d in Document.query.filter(~Document.id.in_(has_sheets)).all()
        if is_spreadsheet(Path(d.stored_filename).suffix)
    ]
    done = 0
    for doc in pending:
        path = Path(current_app.config["UPLOAD_DIR"]) / doc.stored_filename
        if not path.exists():
            continue
        ext = path.suffix.lstrip(".")
        text, ok, sheets = extract_spreadsheet(path, ext)
        if sheets is None:
            continue
        store_sheets(doc, sheets)
        if ok:
            doc.content, doc.extraction_ok = text, True
        db.session.commit()
        done += 1
    if done:
        logger.info("%d planilha(s) estruturada(s) no banco.", done)
    return done


# --------------------------------------------------------------------------- #
# Consulta: linhas que citam os termos + pseudoplanilha paginada
# --------------------------------------------------------------------------- #
SHEET_PAGE_SIZE = 200
MAX_MATCHED_ROWS = 2_000


def mark_cell(value: str, term_re) -> Markup:
    """Celula escapada, com <mark> nos termos."""
    if not value or term_re is None:
        return Markup(escape(value or ""))
    spans = [(m.start(), m.end()) for m in term_re.finditer(value) if m.end() > m.start()]
    if not spans:
        return Markup(escape(value))
    return Markup(render_window(value, 0, len(value), spans))


def match_rows(doc_id: int, fts_query: str, limit: int = MAX_MATCHED_ROWS):
    """Linhas do documento que casam com a expressao FTS5.

    [(sheet_id, row_number, cells_json, text)], na ordem da planilha.
    """
    if not fts_query:
        return []
    return db.session.execute(
        text(
            """
            SELECT r.sheet_id, r.row_number, r.cells, r.text
            FROM sheet_row_fts f JOIN sheet_row r ON r.id = f.rowid
            WHERE sheet_row_fts MATCH :m AND r.document_id = :d
            ORDER BY r.sheet_id, r.row_number
            LIMIT :lim
            """
        ),
        {"m": fts_query, "d": doc_id, "lim": limit},
    ).all()


def row_url(doc_id: int, sheet_id: int, row_number: int) -> str:
    """Link que abre a pseudoplanilha ja na linha (fonte exata)."""
    return url_for(
        "documents.detail", doc_id=doc_id, sheet=sheet_id, row=row_number,
        _anchor=f"r{row_number}",
    )


def render_row_groups(doc, rows, term_re, *, limit_per_sheet: int = 10, filter_q: str = "") -> str:
    """HTML das linhas encontradas, uma aba de navegacao por aba da planilha.

    `rows`: [(sheet_id, row_number, cells_json)]. Mostra ate `limit_per_sheet`
    linhas por aba; com `filter_q`, cada aba ganha o link "abrir na planilha"
    ja filtrado. Colunas vazias em todas as linhas mostradas sao omitidas.
    """
    sheets = {s.id: s for s in doc.sheets}
    by_sheet: dict[int, list] = {}
    for sheet_id, number, cells_json, *_ in rows:
        by_sheet.setdefault(sheet_id, []).append((number, json.loads(cells_json)))

    groups = []
    for sheet_id, items in by_sheet.items():
        sheet = sheets.get(sheet_id)
        if sheet is None:
            continue
        shown = items[:limit_per_sheet]
        width = max([sheet.n_cols] + [len(c) for _, c in shown])
        keep = [i for i in range(width) if any(i < len(c) and c[i] for _, c in shown)]
        header = sheet.header_list
        letters = [column_letter(sheet.col_offset + i) for i in range(width)]
        groups.append(
            {
                "sheet": sheet,
                "columns": [
                    {"letter": letters[i], "name": mark_cell(header[i] if i < len(header) else "", term_re)}
                    for i in keep
                ],
                "rows": [
                    {
                        "number": number,
                        "url": row_url(doc.id, sheet.id, number),
                        "cells": [mark_cell(c[i] if i < len(c) else "", term_re) for i in keep],
                    }
                    for number, c in shown
                ],
                "total": len(items),
                "open_url": url_for("documents.detail", doc_id=doc.id, sheet=sheet.id,
                                    f=filter_q or None),
            }
        )
    if not groups:
        return ""
    groups.sort(key=lambda g: g["sheet"].ordinal)
    # abre na aba com mais linhas encontradas
    active = max(range(len(groups)), key=lambda i: groups[i]["total"])
    return render_template("_sheet_rows.html", doc=doc, groups=groups, active=active,
                           filter_q=filter_q)


def sheet_page(sheet, *, page: int = 1, filter_q: str = "", target_row: int = 0) -> dict:
    """Uma pagina da pseudoplanilha (com filtro opcional por palavra-chave)."""
    from .locator import compile_terms
    from .models import SheetRow
    from .search import build_match_query, tokenize

    query = SheetRow.query.filter(SheetRow.sheet_id == sheet.id)
    term_re = None
    match = build_match_query(filter_q) if filter_q else ""
    if match:
        term_re = compile_terms(tokenize(filter_q))
        ids = (
            select(literal_column("rowid"))
            .select_from(table("sheet_row_fts"))
            .where(text("sheet_row_fts MATCH :m"))
        )
        query = query.filter(SheetRow.id.in_(ids)).params(m=match)

    total = query.count()
    pages = max(1, (total + SHEET_PAGE_SIZE - 1) // SHEET_PAGE_SIZE)
    if target_row and not match:
        before = query.filter(SheetRow.row_number < target_row).count()
        page = before // SHEET_PAGE_SIZE + 1
    page = min(max(page, 1), pages)

    rows = (
        query.order_by(SheetRow.row_number)
        .offset((page - 1) * SHEET_PAGE_SIZE)
        .limit(SHEET_PAGE_SIZE)
        .all()
    )
    width = sheet.n_cols
    out_rows = []
    for r in rows:
        cells = r.cell_list
        width = max(width, len(cells))
        out_rows.append({"number": r.row_number, "cells": cells})
    for r in out_rows:
        r["cells"] = [mark_cell(c, term_re) for c in r["cells"]] + [""] * (width - len(r["cells"]))

    header = sheet.header_list
    return {
        "rows": out_rows,
        "letters": [column_letter(sheet.col_offset + i) for i in range(width)],
        "header": header + [""] * (width - len(header)),
        "total": total,
        "page": page,
        "pages": pages,
        "first": out_rows[0]["number"] if out_rows else 0,
        "last": out_rows[-1]["number"] if out_rows else 0,
    }
