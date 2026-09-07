"""Extracao das imagens embutidas em PDF, DOCX, PPTX e XLSX.

`extract_images(path, ext)` devolve uma lista de dicts prontos para gravar:
`{data, ext, thumb_data, thumb_ext, content_type, width, height, ref}`.

Imagens minusculas (icones, marcadores, espacadores) e duplicatas dentro do
mesmo documento sao descartadas. Tudo e best-effort: falha nao quebra o upload.
"""

import hashlib
import io
import logging

logger = logging.getLogger(__name__)

MIN_DIM = 40          # ignora imagens menores que isso (px) em qualquer lado
MAX_PER_DOC = 40      # teto de imagens guardadas por documento
THUMB_MAX = 480       # maior lado da miniatura (px)

_EXT_BY_FORMAT = {
    "JPEG": "jpg", "PNG": "png", "GIF": "gif", "WEBP": "webp",
    "BMP": "bmp", "TIFF": "tiff", "MPO": "jpg",
}
_CONTENT_TYPE = {
    "jpg": "image/jpeg", "png": "image/png", "gif": "image/gif",
    "webp": "image/webp", "bmp": "image/bmp", "tiff": "image/tiff",
}
# Formatos que o navegador nao exibe bem: convertidos para PNG.
_CONVERT_TO_PNG = {"bmp", "tiff"}


def _iter_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            images = list(page.images)
        except Exception:  # noqa: BLE001
            continue
        for img in images:
            data = getattr(img, "data", None)
            if data:
                yield data, f"Página {page_number}"


def _iter_docx(path):
    from docx import Document as DocxDocument

    part = DocxDocument(str(path)).part
    for rel in part.rels.values():
        if "image" in rel.reltype and not rel.is_external:
            try:
                yield rel.target_part.blob, "Documento"
            except Exception:  # noqa: BLE001
                continue


def _iter_pptx(path):
    from pptx import Presentation

    prs = Presentation(str(path))
    for slide_number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            image = getattr(shape, "image", None)
            if image is None:
                continue
            try:
                yield image.blob, f"Slide {slide_number}"
            except Exception:  # noqa: BLE001
                continue


def _iter_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    try:
        for ws in wb.worksheets:
            for image in getattr(ws, "_images", []):
                try:
                    ref = image._data()  # noqa: SLF001 - API interna do openpyxl
                    yield ref, f"Planilha {ws.title}"
                except Exception:  # noqa: BLE001
                    continue
    finally:
        wb.close()


_ITERATORS = {
    "pdf": _iter_pdf,
    "docx": _iter_docx,
    "pptx": _iter_pptx,
    "xlsx": _iter_xlsx,
}


def _make_thumb(pil_image):
    from PIL import Image

    thumb = pil_image.copy()
    thumb.thumbnail((THUMB_MAX, THUMB_MAX), Image.LANCZOS)
    has_alpha = thumb.mode in ("RGBA", "LA", "P")
    buf = io.BytesIO()
    if has_alpha:
        thumb.convert("RGBA").save(buf, format="PNG", optimize=True)
        return buf.getvalue(), "png"
    thumb.convert("RGB").save(buf, format="JPEG", quality=82, optimize=True)
    return buf.getvalue(), "jpg"


def extract_images(path, ext: str) -> list[dict]:
    ext = ext.lower().lstrip(".")
    iterator = _ITERATORS.get(ext)
    if not iterator:
        return []

    from PIL import Image

    results: list[dict] = []
    seen: set[str] = set()
    try:
        for blob, ref in iterator(path):
            if not blob or len(results) >= MAX_PER_DOC:
                if len(results) >= MAX_PER_DOC:
                    logger.info("Limite de %d imagens atingido em %s", MAX_PER_DOC, path)
                    break
                continue

            digest = hashlib.sha1(blob).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)

            try:
                pil = Image.open(io.BytesIO(blob))
                pil.load()
            except Exception:  # noqa: BLE001
                continue

            width, height = pil.size
            if width < MIN_DIM or height < MIN_DIM:
                continue

            fmt = (pil.format or "").upper()
            out_ext = _EXT_BY_FORMAT.get(fmt, "png")
            data = blob
            if out_ext in _CONVERT_TO_PNG:
                buf = io.BytesIO()
                pil.convert("RGBA" if "A" in pil.mode else "RGB").save(buf, format="PNG")
                data, out_ext = buf.getvalue(), "png"

            try:
                thumb_data, thumb_ext = _make_thumb(pil)
            except Exception:  # noqa: BLE001
                thumb_data, thumb_ext = data, out_ext

            results.append({
                "data": data,
                "ext": out_ext,
                "content_type": _CONTENT_TYPE.get(out_ext, "application/octet-stream"),
                "thumb_data": thumb_data,
                "thumb_ext": thumb_ext,
                "width": width,
                "height": height,
                "ref": ref,
            })
    except Exception:  # noqa: BLE001 - best effort
        logger.exception("Falha ao extrair imagens de %s", path)

    return results
